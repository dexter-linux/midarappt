import io
import os
import tempfile
import requests
import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from duckduckgo_search import DDGS
from PIL import Image

# Streamlit Page Configuration
st.set_page_config(
    page_title="AI Presentation Generator Agent",
    page_icon="📊",
    layout="wide"
)

st.title("📊 Agentic PowerPoint Presentation Generator")
st.markdown(
    "Turn your precise slide content into a professionally styled PowerPoint presentation "
    "complete with automatically retrieved relevant web images."
)

# Sidebar Configuration
with st.sidebar:
    st.header("⚙️ Configuration")
    st.markdown("---")
    st.markdown("### 💡 Slide Formatting Guide")
    st.markdown(
        """
        * Separate slides using `---`.
        * The **first line** of each section becomes the slide title.
        * **Subsequent lines** become the bullet points.
        """
    )

# Default template content
default_content = """# Introduction to Agentic Workflows
- Autonomous decision-making models
- Leveraging multi-agent coordination
- Real-world software integration and deployment

---

# Architecture & Tech Stack
- Large Language Models as reasoning engines
- Tool use via function calling and web search
- Integration with local and cloud backends

---

# Future Outlook & Conclusion
- Scalability across enterprise systems
- Robust error handling and security guardrails
- Q&A and wrap-up discussion
"""

# Main text area for slide content
slide_text = st.text_area("Enter your exact slide content below:", value=default_content, height=320)

def search_and_download_image(query):
    """Searches DuckDuckGo for a relevant image and returns a temporary local file path."""
    try:
        with DDGS() as ddgs:
            results = list(ddgs.images(query, max_results=1))
            if results:
                image_url = results[0]["image"]
                response = requests.get(image_url, timeout=5)
                if response.status_code == 200:
                    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".jpg")
                    temp_file.write(response.content)
                    temp_file.close()
                    return temp_file.name
    except Exception as e:
        print(f"Image search error for query '{query}': {e}")
    return None

if st.button("🚀 Generate Presentation", type="primary"):
    if not slide_text.strip():
        st.error("Please provide valid slide content before generating.")
    else:
        with st.spinner("Assembling slides and fetching web imagery..."):
            prs = Presentation()
            # Set widescreen 16:9 dimensions
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            
            blank_layout = prs.slide_layouts[6] # Blank slide template
            slides_raw = slide_text.split("---")
            
            for raw_slide in slides_raw:
                lines = [l.strip() for l in raw_slide.strip().split("\n") if l.strip()]
                if not lines:
                    continue
                
                title = lines[0].replace("#", "").strip()
                body_lines = lines[1:]
                
                slide = prs.slides.add_slide(blank_layout)
                
                # Add Slide Title
                title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
                tf = title_box.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = title
                p.font.size = Pt(32)
                p.font.bold = True
                p.font.color.rgb = RGBColor(31, 78, 121) # Professional dark blue
                
                # Fetch relevant image based on title
                img_path = search_and_download_image(title)
                
                # Adjust text box width depending on whether an image was found
                content_width = Inches(7.0) if img_path else Inches(11.7)
                content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), content_width, Inches(5.0))
                ctf = content_box.text_frame
                ctf.word_wrap = True
                
                for i, line in enumerate(body_lines):
                    clean_line = line.lstrip("-*• ").strip()
                    if i == 0 and len(ctf.paragraphs) > 0:
                        p_body = ctf.paragraphs[0]
                    else:
                        p_body = ctf.add_paragraph()
                    p_body.text = f"• {clean_line}"
                    p_body.font.size = Pt(18)
                    p_body.font.color.rgb = RGBColor(50, 50, 50)
                    p_body.space_after = Pt(14)
                
                # Insert image if available
                if img_path and os.path.exists(img_path):
                    try:
                        slide.shapes.add_picture(
                            img_path, 
                            Inches(8.2), 
                            Inches(1.8), 
                            width=Inches(4.3)
                        )
                        os.unlink(img_path)
                    except Exception:
                        pass

            # Save presentation to a memory buffer
            pptx_buffer = io.BytesIO()
            prs.save(pptx_buffer)
            pptx_buffer.seek(0)
            
            st.success("Presentation generated successfully!")
            st.download_button(
                label="📥 Download PowerPoint Presentation (.pptx)",
                data=pptx_buffer,
                file_name="ai_generated_presentation.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )
